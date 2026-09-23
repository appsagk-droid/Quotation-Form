import io
import base64
import json
import re
from urllib.error import HTTPError
from urllib.parse import quote
from urllib.request import Request, urlopen
from datetime import date
from pathlib import Path

import pandas as pd
import streamlit as st
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE_DIR = Path(__file__).parent
TEMPLATE_PATH = BASE_DIR / "sample.docx"
COMPANIES_DIR = BASE_DIR / "companies"
ITEM_LIST_PATH = BASE_DIR / "item_list.json"
LOGO_PATH = BASE_DIR / "logo.jpg"
COMPANIES_DIR.mkdir(exist_ok=True)


def slugify(value):
    value = str(value or "").strip()
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-") or "company"


def sanitize_company_name(value):
    return str(value or "").strip()


def github_settings():
    try:
        token = str(st.secrets.get("GITHUB_TOKEN", "")).strip()
        repository = str(st.secrets.get("GITHUB_REPOSITORY", "")).strip()
        branch = str(st.secrets.get("GITHUB_BRANCH", "main")).strip() or "main"
        data_dir = str(st.secrets.get("GITHUB_DATA_DIR", "companies")).strip().strip("/")
    except Exception:
        return None
    if not token or not repository or "/" not in repository:
        return None
    return token, repository, branch, data_dir


def github_request(method, path, settings, payload=None):
    token, repository, branch, _ = settings
    url = f"https://api.github.com/repos/{repository}/contents/{quote(path, safe='/')}"
    if method == "GET":
        url += f"?ref={quote(branch)}"
    body = json.dumps(payload).encode("utf-8") if payload is not None else None
    request = Request(
        url,
        data=body,
        method=method,
        headers={
            "Accept": "application/vnd.github+json",
            "Authorization": f"Bearer {token}",
            "X-GitHub-Api-Version": "2022-11-28",
            "User-Agent": "quotation-form-streamlit",
            **({"Content-Type": "application/json"} if body else {}),
        },
    )
    with urlopen(request, timeout=15) as response:
        return json.loads(response.read().decode("utf-8"))


def github_file_path(settings, filename):
    return f"{settings[3]}/{filename}"


@st.cache_data(ttl=30, show_spinner=False)
def github_companies(settings):
    listing = github_request("GET", settings[3], settings)
    companies = []
    for entry in listing if isinstance(listing, list) else []:
        if not entry.get("name", "").endswith(".json"):
            continue
        file_data = github_request("GET", entry["path"], settings)
        content = base64.b64decode(file_data.get("content", "")).decode("utf-8")
        company = json.loads(content)
        if isinstance(company, dict) and company.get("name"):
            companies.append(company)
    return companies


def github_write_company(company, settings):
    filename = f"{slugify(company.get('name'))}.json"
    path = github_file_path(settings, filename)
    payload = {
        "message": f"Save company: {company.get('name', '').strip()}",
        "content": base64.b64encode(
            json.dumps(company, ensure_ascii=False, indent=2).encode("utf-8")
        ).decode("ascii"),
        "branch": settings[2],
    }
    try:
        existing = github_request("GET", path, settings)
    except HTTPError as error:
        if error.code != 404:
            raise
    else:
        payload["sha"] = existing["sha"]
    result = github_request("PUT", path, settings, payload)
    github_companies.clear()
    return result


def load_companies():
    settings = github_settings()
    if settings:
        return sorted(github_companies(settings), key=lambda item: company_label(item).casefold())
    companies = []
    for path in COMPANIES_DIR.glob("*.json"):
        try:
            company = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            continue
        if isinstance(company, dict) and company.get("name"):
            companies.append(company)
    return sorted(companies, key=lambda item: company_label(item).casefold())


def save_company(company):
    name = sanitize_company_name(company.get("name"))
    if not name:
        raise ValueError("Company name is required.")
    settings = github_settings()
    if settings:
        github_write_company(company, settings)
        return
    path = COMPANIES_DIR / f"{slugify(name)}.json"
    path.write_text(json.dumps(company, ensure_ascii=False, indent=2), encoding="utf-8")


def delete_company(company):
    name = sanitize_company_name(company.get("name"))
    if not name:
        raise ValueError("Company name is required.")
    settings = github_settings()
    if settings:
        path = github_file_path(settings, f"{slugify(name)}.json")
        file_data = github_request("GET", path, settings)
        github_request(
            "DELETE",
            path,
            settings,
            {"message": f"Delete company: {name}", "sha": file_data["sha"], "branch": settings[2]},
        )
        github_companies.clear()
        return
    for path in COMPANIES_DIR.glob("*.json"):
        try:
            existing = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            continue
        if isinstance(existing, dict) and sanitize_company_name(existing.get("name")) == name:
            path.unlink(missing_ok=True)
            return
    raise FileNotFoundError(f"Company '{name}' not found.")


def load_item_list():
    try:
        items = json.loads(ITEM_LIST_PATH.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return []
    return items if isinstance(items, list) else []


def save_item_list(items):
    ITEM_LIST_PATH.write_text(json.dumps(items, ensure_ascii=False, indent=2), encoding="utf-8")


def company_label(company):
    return company.get("name") or "Untitled company"


def format_reference(company, report_date, suffix):
    client = company_label(company).strip()
    date_text = f"{report_date.day}.{report_date.month}.{report_date.year}"
    return f"AGKSol/{date_text}/QUO/{client}/{str(suffix).strip()}"


def get_default_categories():
    return {
        "A": {"title": "", "rows": []},
        "B": {"title": "", "rows": []},
        "C": {"title": "", "rows": []},
        "D": {"title": "", "rows": []},
    }


def category_label(index):
    label = ""
    while index >= 0:
        index, remainder = divmod(index, 26)
        label = chr(65 + remainder) + label
        index -= 1
    return label


def parse_numeric(value):
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    if not text:
        return 0.0
    match = re.findall(r"[-+]?\d*\.?\d+(?:[eE][-+]?\d+)?", text)
    if not match:
        return 0.0
    try:
        return float(match[0])
    except ValueError:
        return 0.0


def auto_calculate_amounts(raw_rows):
    rows = []
    for row in raw_rows or []:
        if not isinstance(row, dict):
            continue
        description = str(row.get("Description", "") or "").strip()
        qty = str(row.get("Qty", "") or "").strip()
        unit_price = str(row.get("U/price", "") or "").strip()
        amount_input = str(row.get("Amount", "") or "").strip()

        qty_number = parse_numeric(qty)
        unit_price_number = parse_numeric(unit_price)
        amount_value = parse_numeric(amount_input) if amount_input else 0.0

        if qty_number and unit_price_number and (not amount_input or amount_value == 0.0):
            amount_value = qty_number * unit_price_number

        if description or qty or unit_price or amount_input:
            rows.append(
                {
                    "Item": "",
                    "Description": description,
                    "Qty": qty,
                    "U/price": unit_price,
                    "Amount": f"{amount_value:,.2f}" if amount_value else "",
                }
            )
    return rows


def normalize_category_rows(raw_rows):
    return auto_calculate_amounts(raw_rows)


def category_value_to_rows(value):
    if value is None:
        return []
    if isinstance(value, list):
        list_rows = []
        for item in value:
            if isinstance(item, dict):
                list_rows.append({
                    "Item": item.get("Item", ""),
                    "Description": item.get("Description", ""),
                    "Qty": item.get("Qty", ""),
                    "U/price": item.get("U/price", ""),
                    "Amount": item.get("Amount", ""),
                })
            elif isinstance(item, str):
                list_rows.append({
                    "Item": "",
                    "Description": item,
                    "Qty": "",
                    "U/price": "",
                    "Amount": "",
                })
        return list_rows
    if isinstance(value, str):
        return [{"Item": "", "Description": value, "Qty": "", "U/price": "", "Amount": ""}]
    return []


def build_item_table(rows):
    cleaned = normalize_category_rows(rows)
    return pd.DataFrame(cleaned, columns=["Item", "Description", "Qty", "U/price", "Amount"])


def item_options(item_list, rows):
    options = [""] + [str(item.get("Item", "")).strip() for item in item_list if item.get("Item")]
    for row in rows:
        item = str(row.get("Item", "") or "").strip()
        if item and item not in options:
            options.append(item)
    return options


def sum_category_amounts(categories):
    total = 0.0
    for section in categories.values():
        rows = section.get("rows", []) if isinstance(section, dict) else section
        for row in normalize_category_rows(rows):
            amount = str(row.get("Amount", "") or "").strip()
            total += parse_numeric(amount)
    return total


def integer_to_words(value):
    value = int(float(value))
    if value == 0:
        return "Zero"

    ones = {
        0: "Zero", 1: "One", 2: "Two", 3: "Three", 4: "Four", 5: "Five",
        6: "Six", 7: "Seven", 8: "Eight", 9: "Nine",
    }
    teens = {
        10: "Ten", 11: "Eleven", 12: "Twelve", 13: "Thirteen", 14: "Fourteen",
        15: "Fifteen", 16: "Sixteen", 17: "Seventeen", 18: "Eighteen", 19: "Nineteen",
    }
    tens = {
        20: "Twenty", 30: "Thirty", 40: "Forty", 50: "Fifty", 60: "Sixty",
        70: "Seventy", 80: "Eighty", 90: "Ninety",
    }

    def under_1000(n):
        if n < 10:
            return ones[n]
        if n < 20:
            return teens[n]
        if n < 100:
            tens_val = n // 10 * 10
            remainder = n % 10
            return tens[tens_val] + (" " + ones[remainder] if remainder else "")
        hundreds = n // 100
        remainder = n % 100
        text = ones[hundreds] + " Hundred"
        if remainder:
            text += " " + under_1000(remainder)
        return text

    if value < 1000:
        return under_1000(value)
    if value < 1000000:
        thousands = value // 1000
        remainder = value % 1000
        text = integer_to_words(thousands) + " Thousand"
        if remainder:
            text += " " + under_1000(remainder)
        return text
    if value < 1000000000:
        millions = value // 1000000
        remainder = value % 1000000
        text = integer_to_words(millions) + " Million"
        if remainder:
            text += " " + integer_to_words(remainder)
        return text
    billions = value // 1000000000
    remainder = value % 1000000000
    text = integer_to_words(billions) + " Billion"
    if remainder:
        text += " " + integer_to_words(remainder)
    return text


def format_written_total(amount):
    total_value = int(round(float(amount)))
    words = integer_to_words(total_value)
    return f"Ringgit Malaysia : {words} Only"


def add_simple_table_borders(table):
    borders = table._tbl.tblPr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table._tbl.tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "B7B7B7")


def remove_paragraph(paragraph):
    paragraph._element.getparent().remove(paragraph._element)


def add_paragraph_after(block, text):
    paragraph = block.part.document.add_paragraph(text)
    block._element.addnext(paragraph._element)
    return paragraph


def style_paragraph(paragraph, size=9, bold=False, space_before=0, space_after=0, line_spacing=1.0):
    paragraph.paragraph_format.space_before = Pt(space_before)
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = line_spacing
    for run in paragraph.runs:
        run.font.name = "Arial"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        run.font.size = Pt(size)
        run.bold = bold if bold else run.bold


def keep_paragraphs_together(paragraphs):
    for index, paragraph in enumerate(paragraphs):
        paragraph.paragraph_format.keep_with_next = index < len(paragraphs) - 1
        paragraph.paragraph_format.keep_together = True


def style_cell(cell, size=9, bold=False, align=None):
    for paragraph in cell.paragraphs:
        if align is not None:
            paragraph.alignment = align
        style_paragraph(paragraph, size=size, bold=bold)


def set_table_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)


def get_usable_width(section):
    return (section.page_width - section.left_margin - section.right_margin) / Inches(1)


def set_table_cell_margins(table, top=0, start=40, bottom=0, end=40):
    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_mar = tc_pr.first_child_found_in("w:tcMar")
            if tc_mar is None:
                tc_mar = OxmlElement("w:tcMar")
                tc_pr.append(tc_mar)
            for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
                margin = tc_mar.find(qn(f"w:{side}"))
                if margin is None:
                    margin = OxmlElement(f"w:{side}")
                    tc_mar.append(margin)
                margin.set(qn("w:w"), str(value))
                margin.set(qn("w:type"), "dxa")


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "nil")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def generate_docx(company, report_date, categories, reference, quotation_title):
    if not TEMPLATE_PATH.is_file():
        raise FileNotFoundError(
            "The DOCX template is missing. Commit sample.docx to the repository "
            "and redeploy the Streamlit app."
        )
    doc = Document(TEMPLATE_PATH)
    for section in doc.sections:
        section.top_margin = Inches(0.18)
        section.bottom_margin = Inches(0.35)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    usable_width = get_usable_width(doc.sections[0])

    company_name = sanitize_company_name(company.get("name"))
    address_lines = [line.strip() for line in str(company.get("address", "")).splitlines() if line.strip()]

    table0 = doc.tables[0]
    client_cell = table0.rows[7].cells[0]
    client_cell.text = ""
    client_paragraph = client_cell.paragraphs[0]
    client_run = client_paragraph.add_run(company_name)
    client_run.bold = True

    logo_cell = table0.rows[0].cells[0]
    logo_cell.text = ""
    header_table = logo_cell.add_table(rows=1, cols=2)
    header_table.autofit = False
    set_table_widths(header_table, [1.45, usable_width - 1.45])
    set_table_cell_margins(header_table)
    remove_table_borders(header_table)
    logo_column, sender_column = header_table.rows[0].cells
    if LOGO_PATH.exists():
        logo_paragraph = logo_column.paragraphs[0]
        logo_paragraph.paragraph_format.space_after = Pt(0)
        logo_paragraph.add_run().add_picture(str(LOGO_PATH), width=Inches(1.25))
    sender_column.text = ""
    sender_lines = [
        ("AGK SOLUTIONS", True),
        ("No.20A, Jalan Tiara 2, Tiara Square Business Centre,", True),
        ("Taman Perindustrian SIME UEP, 47620 Subang Jaya, Selangor", True),
        ("D. E", True),
        ("Tel: 03-8023 9303", True),
        ("email: agksolutions28@gmail.com", True),
    ]
    for index, (text, bold) in enumerate(sender_lines):
        paragraph = sender_column.paragraphs[0] if index == 0 else sender_column.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(text)
        run.bold = bold
        style_paragraph(paragraph, size=8.5, bold=bold, line_spacing=0.95)

    for row in table0.rows[1:5]:
        table0._element.remove(row._element)

    for index, line in enumerate(address_lines[:2]):
        row_index = 4 + index
        if row_index < len(table0.rows):
            table0.rows[row_index].cells[0].text = line
    for row_index in range(4 + max(len(address_lines), 2), 6):
        if row_index < len(table0.rows):
            table0.rows[row_index].cells[0].text = ""

    formatted_date = report_date.strftime("%d %B %Y")
    if len(table0.rows) > 2:
        table0.rows[2].cells[0].text = f"Date : {formatted_date}"
    if len(table0.rows) > 1:
        table0.rows[1].cells[0].text = f"Our Ref: {reference}"
    if len(table0.rows) > 6:
        attention_name = str(company.get("attn", "") or "").strip()
        table0.rows[6].cells[0].text = f"Attn : {attention_name}" if attention_name else ""
    if len(table0.rows) > 8:
        table0.rows[8].cells[0].text = quotation_title.strip()

    for row_index, row in enumerate(table0.rows):
        for cell in row.cells:
            cell.margin_top = Inches(0.02)
            cell.margin_bottom = Inches(0.02)
            style_cell(cell, size=9.5 if row_index in (0, 3, 8) else 9)
    style_cell(table0.rows[0].cells[0], size=10, bold=True)
    style_cell(table0.rows[3].cells[0], size=10, bold=True)
    style_cell(table0.rows[8].cells[0], size=10, bold=True)

    for table in doc.tables[1:]:
        table._element.getparent().remove(table._element)

    intro_paragraph = doc.paragraphs[0]
    for paragraph in list(doc.paragraphs):
        if paragraph._p is not intro_paragraph._p:
            remove_paragraph(paragraph)

    table1 = doc.add_table(rows=1, cols=5)
    add_simple_table_borders(table1)
    table_widths = [usable_width * 0.085, usable_width * 0.515, usable_width * 0.13,
                    usable_width * 0.135, usable_width * 0.135]
    intro_paragraph._p.addnext(table1._element)
    header = table1.rows[0].cells
    for cell, text in zip(header, ["Item", "Description", "Qty", "Price / Unit", "Amount"]):
        cell.text = text
        style_cell(cell, size=8, bold=True, align=1)
        set_cell_shading(cell, "D9EAF7")

    for section_letter, section in categories.items():
        section_rows = section.get("rows", []) if isinstance(section, dict) else section
        rows = normalize_category_rows(section_rows)
        if not rows:
            continue

        section_title = str((section.get("title") if isinstance(section, dict) else "") or "").strip()
        row = table1.add_row()
        row.cells[0].text = ""
        merged_cell = row.cells[1].merge(row.cells[-1])
        merged_cell.text = f"{section_letter}. {section_title or f'Section {section_letter}'}"
        style_cell(merged_cell, size=8, bold=True)

        for idx, item in enumerate(rows, start=1):
            row = table1.add_row()
            item_number = str(item.get("Item", "") or "").strip() or f"{idx}."
            description = str(item.get("Description", "") or "")
            qty = str(item.get("Qty", "") or "")
            unit_price = str(item.get("U/price", "") or "")
            amount = str(item.get("Amount", "") or "")

            row.cells[0].text = item_number
            row.cells[1].text = description
            row.cells[2].text = qty
            row.cells[3].text = unit_price
            row.cells[4].text = amount
            style_cell(row.cells[0], size=8, align=1)
            style_cell(row.cells[1], size=8)
            style_cell(row.cells[2], size=8, align=1)
            style_cell(row.cells[3], size=8, align=1)
            style_cell(row.cells[4], size=8, align=1)

    grand_total = sum_category_amounts(categories)
    total_row = table1.add_row()
    total_row.cells[0].merge(total_row.cells[3]).text = "Grand Total"
    total_row.cells[4].text = f"{grand_total:,.2f}"
    style_cell(total_row.cells[0], size=8, bold=True, align=2)
    style_cell(total_row.cells[4], size=8, bold=True, align=1)
    set_table_widths(table1, table_widths)
    set_table_cell_margins(table1, top=40, start=80, bottom=40, end=80)
    written_total_paragraph = add_paragraph_after(
        table1,
        format_written_total(grand_total),
    )
    signature_paragraph = add_paragraph_after(
        written_total_paragraph,
        "This document is electronically generated. No Signature required.",
    )
    style_paragraph(intro_paragraph, size=9, space_before=2, space_after=4)
    style_paragraph(written_total_paragraph, size=9, bold=True, space_before=1)
    style_paragraph(signature_paragraph, size=9, space_before=2)
    keep_paragraphs_together([written_total_paragraph, signature_paragraph])

    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


st.set_page_config(page_title="Quotation Form", layout="centered")

if "screen" not in st.session_state:
    st.session_state.screen = "home"


def reset_form():
    for key in [
        "selected_company",
        "generated_docx",
        "generated_date",
        "generated_reference",
        "share_text",
        "reference_suffix",
        "report_date",
        "quotation_title",
        "category_count",
        "report_draft",
    ]:
        st.session_state.pop(key, None)
    for key in list(st.session_state):
        if key.startswith("category_title_") or key.startswith("category_"):
            st.session_state.pop(key, None)


if st.session_state.screen == "home":
    st.title("Quotation Generator")
    companies = load_companies()

    if companies:
        st.subheader("Select a company")
        labels = [company_label(company) for company in companies]
        selected_index = st.radio("Companies", labels, index=0, label_visibility="collapsed")
        selected_company = companies[labels.index(selected_index)]

        col1, col2, col3, col4 = st.columns(4)
        with col1:
            if st.button("Continue", type="primary"):
                st.session_state.selected_company = selected_company
                st.session_state.screen = "report"
                st.rerun()
        with col2:
            if st.button("Create new company"):
                st.session_state.screen = "create_company"
                st.rerun()
        with col3:
            if st.button("Edit company"):
                st.session_state.edit_company = selected_company
                st.session_state.screen = "edit_company"
                st.rerun()
        with col4:
            if st.button("Delete company", type="secondary"):
                try:
                    delete_company(selected_company)
                    st.session_state.pop("selected_company", None)
                    st.session_state.pop("edit_company", None)
                    st.success(f"Deleted company: {company_label(selected_company)}")
                    st.rerun()
                except (FileNotFoundError, ValueError) as exc:
                    st.error(str(exc))
    else:
        st.info("No company has been created yet.")
        if st.button("Create new company", type="primary"):
            st.session_state.screen = "create_company"
            st.rerun()

elif st.session_state.screen == "create_company":
    st.title("Create company")
    if st.button("Back"):
        st.session_state.screen = "home"
        st.rerun()

    with st.form("new_company_form"):
        company_name = st.text_input("Company name")
        company_address = st.text_area("Address")
        attention_name = st.text_input("Attention name")
        company_quotation_title = st.text_area(
            "Default quotation title",
            help="This will be used as the starting title when creating a quotation for this company.",
        )
        submitted = st.form_submit_button("Save company", type="primary")

    if submitted:
        if not company_name.strip():
            st.error("Company name is required.")
        else:
            new_company = {
                "name": company_name.strip(),
                "address": company_address.strip(),
                "attn": attention_name.strip(),
                "quotation_title": company_quotation_title.strip(),
            }
            save_company(new_company)
            st.session_state.selected_company = new_company
            st.session_state.screen = "report"
            st.rerun()

elif st.session_state.screen == "edit_company":
    company = st.session_state.edit_company
    st.title("Edit company")
    if st.button("Back"):
        st.session_state.screen = "home"
        st.rerun()

    with st.form("edit_company_form"):
        company_name = st.text_input("Company name", value=str(company.get("name", "")))
        company_address = st.text_area("Address", value=str(company.get("address", "")))
        attention_name = st.text_input("Attention name", value=str(company.get("attn", "")))
        company_quotation_title = st.text_area(
            "Default quotation title",
            value=str(company.get("quotation_title", "")),
            help="This will be used as the starting title when creating a quotation for this company.",
        )
        submitted = st.form_submit_button("Save changes", type="primary")

    if submitted:
        if not company_name.strip():
            st.error("Company name is required.")
        else:
            updated_company = {
                "name": company_name.strip(),
                "address": company_address.strip(),
                "attn": attention_name.strip(),
                "quotation_title": company_quotation_title.strip(),
            }
            save_company(updated_company)
            st.session_state.selected_company = updated_company
            st.session_state.screen = "report"
            st.rerun()

elif st.session_state.screen == "report":
    company = st.session_state.selected_company
    st.title("Quotation")
    st.caption(f"Client: {company_label(company)}")
    item_list = load_item_list()
    draft = st.session_state.get("report_draft", {})
    draft_categories = draft.get("categories", {})
    if "category_count" not in st.session_state:
        st.session_state.category_count = draft.get("category_count", 4)
    if "report_date" not in st.session_state:
        st.session_state.report_date = draft.get("report_date", date.today())
    if "reference_suffix" not in st.session_state:
        st.session_state.reference_suffix = draft.get("reference_suffix", "01")
    if "quotation_title" not in st.session_state:
        default_title = str(company.get("quotation_title", "") or "").strip()
        if not default_title:
            default_title = "MAINTENANCE SERVICES OF 11KV SUBSTATION ELECTRICAL EQUIPMENT AT TOP GLOVE SDN BHD (F 24)"
        st.session_state.quotation_title = draft.get("quotation_title", default_title)

    report_date = st.date_input("Date", key="report_date")
    reference_suffix = st.text_input("Reference number", key="reference_suffix")
    reference = format_reference(company, report_date, reference_suffix)
    st.caption(f"Our Ref: {reference}")
    quotation_title = st.text_area(
        "Quotation title",
        height=100,
        key="quotation_title",
    )

    categories = {}
    for category_index in range(st.session_state.category_count):
        letter = category_label(category_index)
        current = draft_categories.get(letter, {"title": "", "rows": []})
        title_key = f"category_title_{letter}"
        if title_key not in st.session_state:
            st.session_state[title_key] = str(current.get("title", ""))
        category_title = st.text_input(f"Category {letter} name", key=title_key)
        current_rows = current.get("rows", [])
        table_data = build_item_table(current_rows)[["Description", "Qty", "U/price"]]
        if table_data.empty:
            table_data = pd.DataFrame([{"Description": "", "Qty": "", "U/price": ""}])
        editor_key = f"category_{letter}"
        rows = st.data_editor(
            table_data,
            hide_index=True,
            width="stretch",
            num_rows="dynamic",
            column_config={
                "Description": st.column_config.TextColumn("Description", width="large"),
                "Qty": st.column_config.TextColumn("Qty", width="small"),
                "U/price": st.column_config.TextColumn("Unit / price", width="small"),
            },
            key=editor_key,
        )
        records = rows.to_dict("records")
        for idx, record in enumerate(records):
            qty_value = parse_numeric(record.get("Qty"))
            unit_value = parse_numeric(record.get("U/price"))
            amount_value = qty_value * unit_value
            record["Amount"] = f"{amount_value:,.2f}" if amount_value else ""
            records[idx] = record
        categories[letter] = {"title": category_title, "rows": normalize_category_rows(records)}

    if st.button("Add category"):
        st.session_state.category_count += 1
        st.rerun()

    if st.button("Generate document", type="primary"):
        if not reference_suffix.strip():
            st.error("Reference number is required.")
            st.stop()
        if not quotation_title.strip():
            st.error("Quotation title is required.")
            st.stop()
        st.session_state.report_draft = {
            "report_date": report_date,
            "reference_suffix": reference_suffix,
            "quotation_title": quotation_title,
            "category_count": st.session_state.category_count,
            "categories": categories,
        }
        doc_bytes = generate_docx(company, report_date, categories, reference, quotation_title)
        st.session_state.generated_docx = doc_bytes
        st.session_state.generated_date = report_date
        st.session_state.generated_reference = reference
        st.session_state.share_text = (
            f"Quotation\n"
            f"Client: {company_label(company)}\n"
            f"Date: {report_date.strftime('%d %B %Y')}\n"
            f"Our Ref: {reference}\n"
            f"Address: {company.get('address','')}"
        )
        st.session_state.screen = "download"
        st.rerun()

    if st.button("Change company"):
        reset_form()
        st.session_state.screen = "home"
        st.rerun()

elif st.session_state.screen == "download":
    company = st.session_state.selected_company
    st.title("Document ready")
    st.success(f"Quotation for {company_label(company)} has been generated.")

    filename = f"{slugify(company_label(company))}_{st.session_state.generated_date:%Y-%m-%d}"
    st.download_button(
        label="Download Word (.docx)",
        data=st.session_state.generated_docx,
        file_name=f"{filename}.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        width="stretch",
    )

    if st.button("Back to edit quotation"):
        st.session_state.screen = "report"
        st.rerun()

    if st.button("Create another quotation"):
        reset_form()
        st.session_state.screen = "home"
        st.rerun()
